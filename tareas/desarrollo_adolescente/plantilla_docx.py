"""Plantilla Word: hoja 1 horizontal con el mapa conceptual y el resto vertical en blanco."""
from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

AQUI = Path(__file__).resolve().parent
doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)

hoja1 = doc.sections[0]
hoja1.orientation = WD_ORIENT.LANDSCAPE
hoja1.page_width, hoja1.page_height = Cm(29.7), Cm(21)
hoja1.left_margin = hoja1.right_margin = Cm(1.5)
hoja1.top_margin = hoja1.bottom_margin = Cm(1.5)
doc.add_picture(str(AQUI / "mapa_conceptual.png"), width=Cm(26.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

resto = doc.add_section(WD_SECTION.NEW_PAGE)
resto.orientation = WD_ORIENT.PORTRAIT
resto.page_width, resto.page_height = Cm(21), Cm(29.7)
resto.left_margin = resto.right_margin = Cm(2.5)
resto.top_margin = resto.bottom_margin = Cm(2.5)
doc.add_paragraph("")

doc.save(AQUI / "trabajo_plantilla.docx")
