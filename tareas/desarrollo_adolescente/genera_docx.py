"""Genera el trabajo «Desarrollo del adolescente y tecnologías digitales» en .docx.

Uso:  python3 genera_docx.py   (crea desarrollo_adolescente.docx junto a este script)
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
ACENTO = RGBColor(0x22, 0x30, 0x3F)

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(5)
st.paragraph_format.line_spacing = 1.12
for nivel, tam in ((1, 14), (2, 12)):
    h = doc.styles[f"Heading {nivel}"]
    h.font.name = "Calibri"
    h.font.size = Pt(tam)
    h.font.color.rgb = ACENTO
    h.paragraph_format.space_before = Pt(10 if nivel == 1 else 7)
    h.paragraph_format.space_after = Pt(4)

sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21)
sec.left_margin = sec.right_margin = Cm(1.5)
sec.top_margin = sec.bottom_margin = Cm(1)


def p(texto, negrita_inicial=None, justif=True):
    par = doc.add_paragraph()
    if negrita_inicial:
        par.add_run(negrita_inicial).bold = True
    par.add_run(texto)
    if justif:
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return par


def vineta(texto, negrita_inicial=None):
    par = doc.add_paragraph(style="List Bullet")
    if negrita_inicial:
        par.add_run(negrita_inicial).bold = True
    par.add_run(texto)
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return par


# ---------------------------------------------------------------- portada breve
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("El desarrollo del adolescente y la influencia de las tecnologías digitales")
r.bold, r.font.size, r.font.color.rgb = True, Pt(17), ACENTO
s = doc.add_paragraph()
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
s.add_run("Mapa conceptual y explicación  ·  Máster en Profesorado de Educación Secundaria\n"
          "Alumna/o: ______________________    Fecha: ____________").italic = True

# ---------------------------------------------------------------- mapa (horizontal)
doc.add_picture(str(AQUI / "mapa_conceptual.png"), width=Cm(18))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

texto = doc.add_section()
texto.orientation = WD_ORIENT.PORTRAIT
texto.page_width, texto.page_height = Cm(21), Cm(29.7)
texto.left_margin = texto.right_margin = Cm(2.3)
texto.top_margin = texto.bottom_margin = Cm(2)

doc.add_heading("1. Introducción", 1)
p("La adolescencia es el periodo de la vida en el que el niño se convierte en adulto a través de intensos "
  "cambios físicos, psíquicos y sociales (Casas Rivero y Ceñal González-Fierro, 2005). Coincide casi por "
  "completo con la etapa de Educación Secundaria, por lo que conocerla es imprescindible para que el "
  "profesorado comprenda a su alumnado, distinga lo que es una variación normal de lo que es una señal de "
  "alarma, y oriente su práctica. Hoy, además, esta etapa se vive en un entorno digital permanente: el 98 % "
  "de los adolescentes españoles tiene Internet en casa y casi uno de cada tres lo usa de forma habitual "
  "desde antes de los 10 años (Save the Children, 2024). El mapa conceptual de la página anterior organiza "
  "el desarrollo adolescente en sus tres dimensiones —bio-física, psico-afectiva y socio-moral—, sitúa debajo "
  "de cada una la influencia de las tecnologías digitales y termina con nuestra posición como futuros "
  "docentes. Las secciones posteriores explican el mapa de arriba abajo.")


# ---------------------------------------------------------------- explicación
doc.add_heading("2. Características del desarrollo adolescente", 1)
p("El nodo central del mapa define la adolescencia como un tránsito de unos diez años que suele dividirse en "
  "tres fases: temprana (11-13 años), media (14-17) y tardía (17-21). Aunque se describen por separado, son "
  "un continuo; los adolescentes no forman un grupo homogéneo y los aspectos biológicos, intelectuales, "
  "emocionales y sociales no maduran al mismo ritmo, con posibles retrocesos en momentos de estrés. En el aula "
  "de 1.º de ESO conviven, por tanto, alumnos que aún parecen niños con otros ya plenamente puberales. A lo "
  "largo de estas fases el adolescente debe lograr cuatro tareas: independizarse de los padres, adaptarse al "
  "grupo, aceptar su nueva imagen corporal y establecer su identidad sexual, moral y vocacional.")

doc.add_heading("2.1. Dimensión bio-física", 2)
p("Es la base del resto de cambios. La pubertad se inicia cuando el hipotálamo aumenta la secreción pulsátil "
  "de GnRH, que estimula en la hipófisis la liberación de FSH y LH y, con ellas, la producción gonadal de "
  "estrógenos y andrógenos. El rango de normalidad es amplio: el 95 % de las chicas inicia la pubertad entre "
  "los 8,5 y los 13 años (menarquia media a los 12,4) y el 95 % de los chicos entre los 9,5 y los 14; además, "
  "la edad de inicio se ha adelantado unos 3-4 meses por década. En la adolescencia temprana domina el "
  "crecimiento rápido y la aparición de los caracteres sexuales secundarios (estadios de Tanner); en la media "
  "se alcanza alrededor del 95 % de la talla adulta y en la tardía la madurez física es completa.")
p("Estos cambios hacen perder la imagen corporal infantil, lo que genera una gran preocupación y curiosidad "
  "por el propio cuerpo y una comparación constante con los iguales. Junto a ello, el cerebro vive su segundo "
  "gran periodo de reorganización tras los 0-3 años: se forman muchas conexiones y su plasticidad lo hace "
  "especialmente sensible a los estímulos del entorno, al sueño y a los hábitos.")

doc.add_heading("2.2. Dimensión psico-afectiva", 2)
p("Comprende el pensamiento, las emociones y la identidad. Cognitivamente se pasa del pensamiento concreto "
  "(en la fase temprana no se perciben las consecuencias futuras de los actos) al pensamiento abstracto: en "
  "la fase media el adolescente disfruta discutiendo ideas y se interesa por temas idealistas, aunque con el "
  "estrés vuelve a razonar en concreto; en la tardía está orientado al futuro. Afectivamente, el adolescente "
  "temprano se siente centro de una «audiencia imaginaria» que lo observa, lo que explica su extremo sentido "
  "del ridículo y su egocentrismo. En la fase media aparece una sensación de omnipotencia e invulnerabilidad "
  "(«a mí no me va a pasar») que favorece las conductas de riesgo: alcohol, tabaco, drogas o relaciones "
  "sexuales sin protección. Casas y Ceñal recuerdan que muchas de estas conductas responden a una necesidad "
  "sana de retos y de sentirse competente: cuando el entorno no ofrece desafíos (deporte, arte, música, "
  "responsabilidades), el adolescente se los inventa, a veces fuera de la norma. Todo ello se integra en la "
  "gran tarea de la etapa, la construcción de una identidad y una autoestima propias.")

doc.add_heading("2.3. Dimensión socio-moral", 2)
p("Abarca las relaciones y la construcción de valores. El adolescente lucha por emanciparse y controlar su "
  "vida, lo que genera conflicto con la familia; sin embargo, los padres siguen siendo necesarios como "
  "referencia estable («los padres permanecen, el grupo cambia»). El grupo de iguales, primero del mismo "
  "sexo, adquiere un peso enorme: dicta la forma de vestir, hablar y comportarse y sirve para afirmar la "
  "autoimagen y elaborar un código de conducta propio. La necesidad de pertenencia es tan alta que algunos "
  "jóvenes prefieren integrarse en grupos marginales antes que quedarse solos. En la fase tardía el grupo "
  "pierde peso en favor de amistades individuales y relaciones de pareja estables, recíprocas y capaces de "
  "proyectar un futuro común; las relaciones familiares pasan a ser de adulto a adulto. En paralelo se "
  "consolida el juicio moral autónomo y la responsabilidad.")

doc.add_heading("3. Influencia de las tecnologías digitales en cada dimensión", 1)
p("La parte inferior de cada columna del mapa (recuadros discontinuos) recoge cómo influyen las tecnologías "
  "digitales. Las fuentes manejadas coinciden en que el problema no es la tecnología en sí, sino su uso "
  "masivo, temprano y sin acompañamiento: más del 90 % de los adolescentes se conecta al menos una hora "
  "diaria fuera de las tareas escolares, el 21 % dice estar «permanentemente conectado» y el 16,5 % supera "
  "las cuatro horas al día, cifra que se duplica entre quienes empezaron antes de los 10 años (Save the "
  "Children, 2024).")

doc.add_heading("3.1. Sobre la dimensión bio-física", 2)
p("El tiempo de pantalla desplaza actividades básicas para el desarrollo físico: el 18 % de los adolescentes "
  "reconoce que duerme menos por estar conectado y el 12 % que deja de hacer deporte o extraescolares. Los "
  "profesionales hablan de un «ocio paralizante» que sustituye el movimiento y el contacto físico. Desmurget "
  "(2020) añade que la falta de sueño deteriora directamente la atención diurna y que las pantallas lúdicas "
  "actúan sobre el sistema de recompensa (pequeñas descargas de dopamina con cada notificación o «me gusta»): "
  "un metaanálisis con más de 150.000 menores relaciona el consumo de pantallas con el déficit de atención, y "
  "una o tres horas diarias de televisión a los 14 años multiplican por 1,4 el riesgo de problemas de atención "
  "a los 16 (por 2,9 si se superan las tres horas). Por último, las redes exponen el cuerpo adolescente, recién "
  "transformado y fuente de inseguridad, a cánones de belleza inalcanzables, a la sexualización y a la "
  "cosificación, con especial impacto en las chicas.")

doc.add_heading("3.2. Sobre la dimensión psico-afectiva", 2)
p("Save the Children encuentra una asociación clara entre horas de conexión y malestar emocional: entre "
  "quienes se conectan más de cuatro horas diarias, el 30,9 % siente siempre o casi siempre que las "
  "dificultades le superan, frente al 18,9 % del resto; las chicas refieren este malestar el doble que los "
  "chicos. No se puede afirmar causalidad, pero sí que el adolescente renuncia a otras actividades (el 27,6 % "
  "lee menos y el 31 % estudia menos) y que su autoestima depende en parte del «like», una validación social "
  "efímera que engancha. Muchos no son conscientes del tiempo que pasan en plataformas como TikTok, cuyo "
  "algoritmo encadena vídeos sin fin.")
p("Desmurget presenta la inteligencia como «primera víctima» porque las pantallas atacan tres pilares del "
  "desarrollo cognitivo: (1) las interacciones humanas, que se reducen en cantidad y calidad (el cerebro aprende "
  "mucho peor de una persona en vídeo que de una presente: «efecto deficitario del vídeo»); (2) el lenguaje, "
  "porque disminuyen el diálogo y, sobre todo, la lectura, única vía para enriquecer el vocabulario y la "
  "sintaxis más allá de un nivel básico; y (3) la concentración, dañada por la multitarea y las interrupciones "
  "constantes —basta tener el móvil al alcance de la mano para rendir menos—. Para el pensamiento abstracto y "
  "la identidad reflexiva que la adolescencia debería consolidar, esta «atención saqueada» es un obstáculo serio.")

doc.add_heading("3.3. Sobre la dimensión socio-moral", 2)
p("Aquí la influencia es ambivalente. Por un lado, el entorno digital permite mantener amistades, encontrar "
  "comunidades afines y ejercer derechos como la información, la expresión o la participación, reconocidos por "
  "la Observación General n.º 25 del Comité de los Derechos del Niño. Por otro, el grupo de iguales —tan decisivo "
  "en esta etapa— se traslada a las redes, donde rigen la presión por la visibilidad, las modas cambiantes y la "
  "imitación de conductas «populares» que pasan del «ser social digital» al físico. La pantalla favorece una "
  "deshumanización de la interacción (se dice lo que nunca se diría cara a cara), base del ciberacoso y los "
  "mensajes de odio: en torno al 20 % ha vivido o presenciado situaciones de posible acoso online. Las redes "
  "también se usan para controlar a la pareja (el 45 % cree que mirar su móvil está bien, al menos a veces), "
  "el 36 % habla con desconocidos y el 58 % se ha topado con pornografía sin buscarla, que sin educación "
  "afectivo-sexual se convierte en modelo de relación. Además, la conexión se vive en soledad física, lo que "
  "dificulta pedir ayuda, y resta tiempo a la familia (17 %), justo cuando las relaciones intrafamiliares siguen "
  "siendo cruciales para el éxito escolar y la prevención de conductas de riesgo (Desmurget, 2020). Llama la "
  "atención que muchos adolescentes no consideren delito enviar fotos sexuales sin permiso (un tercio) o "
  "difundir mensajes de odio (43 %), lo que revela un juicio moral digital aún inmaduro.")

doc.add_heading("3.4. ¿Por qué la adolescencia es tan sensible a lo digital?", 2)
p("El recuadro naranja del mapa sintetiza tres factores. Primero, el diseño persuasivo: las plataformas, "
  "pensadas para adultos, buscan maximizar el tiempo y la atención del usuario, y encuentran en el adolescente "
  "—con baja percepción del riesgo, necesidad de pertenencia y un sistema de recompensa muy activo— un usuario "
  "ideal. Segundo, el desplazamiento: cada hora de pantalla se resta al sueño, la lectura, el estudio, el "
  "deporte o la conversación. Tercero, el mito del «nativo digital»: haber nacido rodeado de tecnología no "
  "enseña a usarla ni a protegerse; son más bien «huérfanos digitales», a quienes nadie ha educado en ese "
  "entorno, y cuyas familias tienden a fiscalizar más que a acompañar.")

doc.add_heading("4. Nuestra posición como futuros profesores", 1)
p("Ni la tecnofilia que introduce pantallas en el aula por su mera novedad ni la prohibición total que "
  "desentiende a la escuela del mundo en que vive su alumnado. Defendemos un uso pedagógico, crítico y "
  "limitado de las tecnologías digitales en los centros, concretado en cinco principios (base del mapa):")
vineta("La evidencia recogida por Desmurget muestra que se aprende mejor de un ser humano presente que de "
       "una pantalla, y que el lenguaje y el pensamiento se desarrollan hablando, leyendo y argumentando. La "
       "relación docente-alumno, el debate en clase y la lectura —especialmente en papel— deben ser el centro.",
       "Primacía de lo humano. ")
vineta("Usaremos herramientas digitales cuando aporten un valor didáctico que no se logre de otra forma "
       "(simulaciones, acceso a fuentes, producción de contenidos, accesibilidad para alumnado con necesidades "
       "específicas), con objetivos claros y tiempos acotados, evitando la multitarea.",
       "Tecnología con propósito. ")
vineta("Apoyamos que el teléfono personal no se use en el aula ni en los recreos de la ESO, mediante normas "
       "claras, explicadas y, en lo posible, consensuadas con el alumnado. Así se protege la atención, se "
       "favorece la interacción cara a cara y se reducen los conflictos de convivencia.",
       "Móvil fuera del aula. ")
vineta("Precisamente porque no son «nativos digitales», la escuela debe enseñar ciudadanía digital: "
       "pensamiento crítico ante la información y los algoritmos, privacidad, prevención del ciberacoso, "
       "empatía online, y educación en igualdad y afectivo-sexual que contrarreste la pornografía y los "
       "estereotipos sexistas, tal como recomienda Save the Children y exige la LOPIVI.",
       "Educación digital. ")
vineta("La prohibición o el control por sí solos no forman usuarios responsables: el 37 % sabe saltarse los "
       "controles. Hay que coordinarse con las familias, dar ejemplo de un uso moderado como adultos, escuchar "
       "al alumnado (enfoque de derechos) y estar atentos a señales de malestar —frustración, aislamiento, "
       "aburrimiento— para derivar al departamento de orientación cuando sea necesario.",
       "Acompañar, no solo prohibir. ")
p("En definitiva, conocer las tres dimensiones del desarrollo adolescente nos permite entender por qué las "
  "tecnologías les afectan tanto y nos obliga a ofrecerles lo que la pantalla no puede dar: retos reales que les "
  "hagan sentirse competentes, relaciones humanas de calidad y criterio para habitar el mundo digital con libertad "
  "y seguridad.")

doc.add_heading("Referencias", 1)
for ref in (
    "Casas Rivero, J. J. y Ceñal González-Fierro, M. J. (2005). Desarrollo del adolescente. Aspectos físicos, "
    "psicológicos y sociales. Pediatría Integral, IX(1), 20-24.",
    "Desmurget, M. (2020). Desarrollo: la inteligencia es la primera víctima. En La fábrica de cretinos "
    "digitales. Los peligros de las pantallas para nuestros hijos (pp. 256-297). Península.",
    "Save the Children España (2024). Derechos sin conexión. Un análisis sobre derechos de la infancia y la "
    "adolescencia y su protección en el entorno digital. Save the Children España.",
):
    par = doc.add_paragraph(ref)
    par.paragraph_format.left_indent = Cm(1)
    par.paragraph_format.first_line_indent = Cm(-1)

salida = AQUI / "desarrollo_adolescente.docx"
doc.save(salida)
print(salida)
