#!/usr/bin/env python3
"""Junta TODOS los análisis ddPCR (E1 = 01-oct, E2 = 08-oct) en un único Excel con gráficas incrustadas.
Uso: python3 genera_excel.py   ->  ddpcr_resultados.xlsx"""
import os, pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
E1, E2 = f"{AQUI}/ddpcr_13", f"{AQUI}/ddpcr_08oct"
rd = lambda base, f: pd.read_csv(f"{base}/resultados/{f}")
HEAD = PatternFill("solid", fgColor="DDE6F0")

def escribe(ws, df, r0=1, titulo=None):
    if titulo:
        ws.cell(r0, 1, titulo).font = Font(bold=True, size=12); r0 += 1
    for j, c in enumerate(df.columns, 1):
        x = ws.cell(r0, j, c); x.font = Font(bold=True); x.fill = HEAD
    for i, row in enumerate(df.itertuples(index=False), r0 + 1):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, None if pd.isna(v) else (round(v, 3) if isinstance(v, float) else v))
    return r0 + len(df) + 2
def ancho(ws, w=16):
    for j in range(1, ws.max_column + 1): ws.column_dimensions[get_column_letter(j)].width = w
def imagen(ws, base, fn, cell, w=520):
    im = Image(f"{base}/figuras/{fn}"); im.width, im.height = w, int(w * im.height / im.width); ws.add_image(im, cell)

wb = Workbook()

# ---------------- Leeme ----------------
ws = wb.active; ws.title = "Leeme"
txt = [
 "Análisis ddPCR · tesis KRT10 · #13 (silenciamiento alelo-específico)", "",
 "HOJAS:  E1 = experimento 01-oct (WT2 y MUT1)  ·  E2 = experimento 08-oct (solo MUT1)  ·  Resumen global = comparación de ambos.", "",
 "MÉTODO COMÚN",
 "1) Por pocillo se conserva el target con más copias/µL (WT2 → T2/VIC; MUT1 → T1/FAM).",
 "2) Se promedian réplicas; se normaliza cada condición contra la media de su control negativo (c-).",
 "3) Se corrige además por el RNA de cada pocillo/réplica. Dos tipos de barra de error: 'SD propagada' (incluye la incertidumbre del c-) y 'SD propia' (cada grupo con las SD de sus réplicas; el c- también lleva su SD).", "",
 "E1 (01-oct): 2 pocillos por condición, tratados como réplicas biológicas. RNA = µL de libreta ÷ 2. A06 (MUT1 #13 1NM) es atípico: se da con y sin él (sin él n = 1, sin SD).",
 "E2 (08-oct): 2 réplicas biológicas por condición, cada una con 2 pocillos técnicos (técnicas → media por biológica → media ± SD de las 2 biológicas).",
 "     RNA = ng/µL de la libreta (el '×2' ya viene aplicado en esos valores). Unidades corregidas: copias/µL ÷ ng/µL de RNA.",
 "     Excluidos: B01, B02 (sin nombre, ~0 copias) y B03 'wt3' (T2 saturado, no es MUT1).", "",
 "AVISOS",
 "· n = 2 por grupo: SD poco fiables y sin test estadístico.",
 "· E2: A09–A12 (#13 0,5NM) tienen < 10 000 gotas (7 600–8 500). La réplica biológica 1 de #13 1NM es baja (707) frente a la 2 (1829): mucha variación entre biológicas.",
 "· E1 y E2 usan unidades de RNA distintas (µL vs ng/µL): compara los 'veces vs c-', no las copias absolutas corregidas."]
for i, t in enumerate(txt, 1): ws.cell(i, 1, t)
ws["A1"].font = Font(bold=True, size=14)
for c in ("A5", "A15"): ws[c].font = Font(bold=True)
ws.column_dimensions["A"].width = 150

# ---------------- Resumen global (MUT1, fold vs c-) ----------------
filas = []
for met, suf in (("sin corregir", ""), ("corregido por RNA", "_porRNA")):
    for sen, nom in (("", "con A06"), ("_sin_A06", "sin A06")):
        n = rd(E1, f"3_normalizado_vs_control{suf}{sen}.csv"); n = n[n.Genotipo == "MUT1"]
        for _, r in n.iterrows():
            filas.append(dict(Experimento=f"E1 (01-oct) · {nom}", Metrica=met, Condicion=r.Condicion, n_bio=int(r.n_muestra),
                              veces_vs_c=r.fold_vs_control, SD_propia=r.SD_propia, SD_propagada=r.SD_propagada))
    n = rd(E2, f"3_normalizado_vs_control{suf}.csv")
    for _, r in n.iterrows():
        filas.append(dict(Experimento="E2 (08-oct)", Metrica=met, Condicion=r.Condicion, n_bio=int(r.n_bio),
                          veces_vs_c=r.fold_vs_control, SD_propia=r.SD_propia, SD_propagada=r.SD_propagada))
g = pd.DataFrame(filas)
g["_o"] = g.Condicion.map({"#13 0,5NM": 0, "#13 1NM": 1})
g = g.sort_values(["Metrica", "_o", "Experimento"]).drop(columns="_o")
g.columns = ["Experimento", "Corrección", "Condición", "n", "Veces vs c-", "SD propia", "SD propagada"]
ws = wb.create_sheet("Resumen global"); escribe(ws, g, 1, "MUT1 · veces vs control negativo (E1 y E2)"); ancho(ws, 24)

# ---------------- E1 ----------------
VAR = [("Sin corregir · con A06", "", ""), ("Sin corregir · sin A06", "", "_sin_A06"),
       ("Corregido por RNA · con A06", "_porRNA", ""), ("Corregido por RNA · sin A06", "_porRNA", "_sin_A06")]
d = rd(E1, "1_target_dominante_por_pocillo.csv").rename(columns={
    "Copias_uL": "Copias/µL (target dominante)", "Conc_otro_target": "Copias/µL otro target", "Pct_otro": "% otro target",
    "Accepted Droplets": "Gotas", "Positives": "Positivas", "RNA_uL": "RNA µL (libreta÷2)", "Copias_por_uL_RNA": "Copias/µL ÷ µL RNA"})
ws = wb.create_sheet("E1 Pocillos"); escribe(ws, d); ancho(ws, 18)
ws = wb.create_sheet("E1 Medias"); r = 1
for tit, suf, sen in VAR: r = escribe(ws, rd(E1, f"2_medias_replicas{suf}{sen}.csv"), r, tit)
ancho(ws)
ws = wb.create_sheet("E1 Normalizado"); r = 1
REN = {"fold_vs_control": "veces vs c-", "SD_propagada": "SD propagada (incluye c-)", "SD_propia": "SD propia del grupo", "SD_control_propia": "SD propia del c-"}
for tit, suf, sen in VAR: r = escribe(ws, rd(E1, f"3_normalizado_vs_control{suf}{sen}.csv").rename(columns=REN), r, tit)
ancho(ws, 18)
c1 = rd(E1, "4_comparacion_con_sin_A06.csv")
piv = c1.assign(col=c1.Metrica + " · " + c1.A06).pivot_table(index=["Genotipo", "Condicion"], columns="col", values="fold", sort=False).reset_index()
ws = wb.create_sheet("E1 Resumen"); escribe(ws, piv, 1, "E1 · veces vs c- (resumen)"); ancho(ws, 26)
ws = wb.create_sheet("E1 Gráficas")
ws["A1"] = "Normalizado vs c- (SD propagada)"; ws["A1"].font = Font(bold=True)
def grid(base, patron, fila0, w=560, paso=22):
    for k, (tit, suf, sen) in enumerate(VAR):
        imagen(ws, base, patron.format(suf=suf, sen=sen), f"{'A' if k % 2 == 0 else 'L'}{fila0 + (k // 2) * paso}", w)
grid(E1, "fig4_normalizado{suf}{sen}.png", 3)
ws["A48"] = "Medias ± SD por réplica"; ws["A48"].font = Font(bold=True); grid(E1, "fig3_medias_replicas{suf}{sen}.png", 50)
ws["A94"] = "Normalizado vs c- con SD propia de cada grupo"; ws["A94"].font = Font(bold=True); grid(E1, "fig4_normalizado_SDpropia{suf}{sen}.png", 96)
ws["A140"] = "MUT1: con vs sin A06"; ws["A140"].font = Font(bold=True); imagen(ws, E1, "fig7_comparacion_con_sin_A06.png", "A142", 760)
ws["A164"] = "Datos crudos y control de calidad"; ws["A164"].font = Font(bold=True)
imagen(ws, E1, "fig1_targets_por_pocillo.png", "A166", 620); imagen(ws, E1, "fig5_gotas.png", "L166", 480); imagen(ws, E1, "fig6_RNA_por_pocillo.png", "A190", 480)

# ---------------- E2 ----------------
d = rd(E2, "1_target_dominante_por_pocillo.csv").rename(columns={
    "Rep_bio": "Réplica biológica", "Target_dominante": "Target", "Copias_uL": "Copias/µL (target dominante)", "Conc_otro_target": "Copias/µL otro target",
    "Pct_otro": "% otro target", "Accepted Droplets": "Gotas", "Positives": "Positivas", "RNA_ngul": "RNA ng/µL (libreta)", "Copias_por_ng_RNA": "Copias/µL ÷ ng/µL RNA"})
ws = wb.create_sheet("E2 Pocillos"); r = escribe(ws, d, 1, "Pocillos usados (técnicas)")
escribe(ws, rd(E2, "0_pocillos_excluidos.csv"), r, "Pocillos excluidos"); ancho(ws, 18)
ws = wb.create_sheet("E2 Medias"); r = 1
for tit, suf in (("Sin corregir", ""), ("Corregido por RNA", "_porRNA")):
    r = escribe(ws, rd(E2, f"2a_replicas_biologicas{suf}.csv"), r, f"{tit} · por réplica biológica (media de las técnicas)")
    r = escribe(ws, rd(E2, f"2b_medias_condicion{suf}.csv"), r, f"{tit} · por condición (media ± SD de las 2 biológicas)")
ancho(ws)
ws = wb.create_sheet("E2 Normalizado"); r = 1
for tit, suf in (("Sin corregir", ""), ("Corregido por RNA", "_porRNA")):
    r = escribe(ws, rd(E2, f"3_normalizado_vs_control{suf}.csv").rename(columns=REN | {"ratios_replicas_bio": "ratios de cada biológica"}), r, tit)
ancho(ws, 20)
ws = wb.create_sheet("E2 Gráficas")
cfg = [("Normalizado vs c- · SD propagada (incluye c-)", "fig7_normalizado_SDpropagada{}.png"),
       ("Normalizado vs c- · SD propia de cada grupo", "fig7_normalizado_SDpropia{}.png"),
       ("Medias ± SD de las 2 réplicas biológicas", "fig4_medias_condicion{}.png"),
       ("Técnicas y media por réplica biológica", "fig3_tecnicas_y_biologicas{}.png")]
fila = 1
for tit, pat in cfg:
    ws.cell(fila, 1, tit + "   (izquierda: sin corregir · derecha: corregido por RNA)").font = Font(bold=True)
    imagen(ws, E2, pat.format(""), f"A{fila + 1}", 500); imagen(ws, E2, pat.format("_porRNA"), f"L{fila + 1}", 500); fila += 25
ws.cell(fila, 1, "Datos crudos y control de calidad").font = Font(bold=True)
imagen(ws, E2, "fig1_targets_por_pocillo.png", f"A{fila + 1}", 620); imagen(ws, E2, "fig2_target_dominante.png", f"L{fila + 1}", 620)
imagen(ws, E2, "fig5_gotas.png", f"A{fila + 24}", 460); imagen(ws, E2, "fig6_RNA_por_replica.png", f"L{fila + 24}", 460)

wb.save(f"{AQUI}/ddpcr_resultados.xlsx")
