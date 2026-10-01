#!/usr/bin/env python3
"""Junta los resultados de ddpcr_13 en un único Excel (hojas + gráficas incrustadas)."""
import os, pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
R = lambda f: pd.read_csv(f"{AQUI}/resultados/{f}")
VAR = [("Sin corregir · con A06", "", ""), ("Sin corregir · sin A06", "", "_sin_A06"),
       ("Corregido por RNA · con A06", "_porRNA", ""), ("Corregido por RNA · sin A06", "_porRNA", "_sin_A06")]

wb = Workbook()
HEAD = PatternFill("solid", fgColor="DDE6F0")
def escribe(ws, df, r0=1, titulo=None):
    if titulo:
        ws.cell(r0, 1, titulo).font = Font(bold=True, size=12); r0 += 1
    for j, c in enumerate(df.columns, 1):
        x = ws.cell(r0, j, c); x.font = Font(bold=True); x.fill = HEAD
    for i, row in enumerate(df.itertuples(index=False), r0 + 1):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, None if pd.isna(v) else (round(v, 3) if isinstance(v, float) else v))
    return r0 + len(df) + 2
def ancho(ws, w=16):
    for j in range(1, ws.max_column + 1): ws.column_dimensions[get_column_letter(j)].width = w

# 1 Léeme
ws = wb.active; ws.title = "Leeme"
txt = ["ddPCR #13 · WT2 / MUT1 · #13 1NM y 0,5NM", "",
 "Pasos:",
 "1) Por pocillo se conserva el target con más copias/µL (WT2 → T2/VIC; MUT1 → T1/FAM). B01 (sin nombre) excluido.",
 "2) Media y SD de las réplicas (2 pocillos por condición).",
 "3) RNA por pocillo = valor de la libreta ÷ 2 (se diluyó después). Corregido = copias/µL ÷ µL de RNA.",
 "4) Normalización: media de cada condición ÷ media de su control negativo (wt2 c- / mut1 c-).",
 "5) Todo se da con A06 y sin A06.", "",
 "Avisos:",
 "· A06 (MUT1 #13 1NM) es atípico: 26,5 copias/µL frente a 1497 en su réplica A05.",
 "· Sin A06, MUT1 1NM tiene n = 1: no hay SD (celda vacía), solo valor puntual.",
 "· n = 2 por grupo en el resto: SD poco fiables, sin test estadístico.",
 "· Los controles varían bastante entre réplicas (mut1 c-: 3987 vs 2374), lo que se traslada a los cocientes.",
 "· Supuesto: los valores de la libreta son µL de RNA y siguen el orden A01–A12.",
 "· Unidades corregidas: copias/µL de reacción por µL de RNA (no se conoce el volumen de reacción)."]
for i, t in enumerate(txt, 1): ws.cell(i, 1, t)
ws["A1"].font = Font(bold=True, size=14); ws["A3"].font = ws["A10"].font = Font(bold=True)
ws.column_dimensions["A"].width = 120

# 2 Pocillos
d = R("1_target_dominante_por_pocillo.csv").rename(columns={
    "Copias_uL": "Copias/µL (target dominante)", "Conc_otro_target": "Copias/µL otro target",
    "Pct_otro": "% otro target", "Accepted Droplets": "Gotas", "Positives": "Positivas",
    "RNA_uL": "RNA µL (libreta÷2)", "Copias_por_uL_RNA": "Copias/µL ÷ µL RNA"})
ws = wb.create_sheet("Pocillos"); escribe(ws, d); ancho(ws, 18)

# 3 Medias
ws = wb.create_sheet("Medias réplicas"); r = 1
for tit, suf, sen in VAR:
    r = escribe(ws, R(f"2_medias_replicas{suf}{sen}.csv"), r, tit)
ancho(ws, 16)

# 4 Normalizado
ws = wb.create_sheet("Normalizado vs c-"); r = 1
for tit, suf, sen in VAR:
    n = R(f"3_normalizado_vs_control{suf}{sen}.csv").rename(columns={"fold_vs_control": "veces vs c-", "SD_propagada": "SD"})
    r = escribe(ws, n, r, tit)
ancho(ws, 18)
# resumen compacto
cmp_ = R("4_comparacion_con_sin_A06.csv")
piv = cmp_.assign(col=cmp_.Metrica + " · " + cmp_.A06).pivot_table(index=["Genotipo", "Condicion"], columns="col", values="fold", sort=False).reset_index()
ws2 = wb.create_sheet("Resumen veces vs c-"); escribe(ws2, piv, 1, "Veces vs control negativo (resumen)"); ancho(ws2, 26)
wb.move_sheet("Resumen veces vs c-", offset=-1)

# 5 Gráficas
ws = wb.create_sheet("Gráficas")
ws["A1"] = "Normalizado vs c- (arriba) y medias con réplicas (abajo): sin corregir | sin corregir sin A06 | por RNA | por RNA sin A06"; ws["A1"].font = Font(bold=True)
def img(fn, cell):
    im = Image(f"{AQUI}/figuras/{fn}"); im.width, im.height = 560, int(560 * im.height / im.width); ws.add_image(im, cell)
for k, (tit, suf, sen) in enumerate(VAR):
    col = "A" if k % 2 == 0 else "L"; row = 3 if k < 2 else 25
    img(f"fig4_normalizado{suf}{sen}.png", f"{col}{row}")
ws["A48"] = "Medias ± SD por réplica"; ws["A48"].font = Font(bold=True)
for k, (tit, suf, sen) in enumerate(VAR):
    col = "A" if k % 2 == 0 else "L"; row = 50 if k < 2 else 72
    img(f"fig3_medias_replicas{suf}{sen}.png", f"{col}{row}")

wb.save(f"{AQUI}/../ddpcr_13_resultados.xlsx")
